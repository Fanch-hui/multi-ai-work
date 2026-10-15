"""Create and modify real Office files with the bundled authoring libraries."""
import json
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.shared import Inches as DocxInches
from pptx import Presentation
from pptx.util import Inches
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as SheetImage
from PIL import Image

root = Path(sys.argv[1]).resolve()
root.mkdir(parents=True, exist_ok=True)
tempfile.tempdir = str(root)
picture = root / 'figure.png'
Image.new('RGB', (80, 40), '#2979ff').save(picture)

doc = Document()
doc.add_heading('中文工作报告', 0)
doc.add_paragraph('实际文档内容')
table = doc.add_table(rows=2, cols=2)
table.cell(0, 0).text, table.cell(0, 1).text = '项目', '结果'
table.cell(1, 0).text, table.cell(1, 1).text = '验证', '完成'
doc.add_picture(str(picture), width=DocxInches(1))
doc_path = root / 'report.docx'
doc.save(doc_path)
edited_doc = Document(doc_path)
edited_doc.add_paragraph('已修改正文')
edited_doc.save(doc_path)
assert any(p.text == '已修改正文' for p in Document(doc_path).paragraphs)

deck = Presentation()
slide = deck.slides.add_slide(deck.slide_layouts[5])
slide.shapes.title.text = '中文工作报告'
slide.shapes.add_picture(str(picture), Inches(.5), Inches(1.5), width=Inches(1))
data = CategoryChartData()
data.categories = ['第一项', '第二项']
data.add_series('完成量', [3, 5])
slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(2), Inches(1.5), Inches(6), Inches(3), data)
ppt_path = root / 'slides.pptx'
deck.save(ppt_path)
edited_deck = Presentation(ppt_path)
edited_deck.slides.add_slide(edited_deck.slide_layouts[5]).shapes.title.text = '已修改第二页'
edited_deck.save(ppt_path)
assert len(Presentation(ppt_path).slides) == 2

book = Workbook()
sheet = book.active
sheet.title = '月度统计'
sheet.append(['中文工作报告', '数值'])
sheet.append(['第一项', 3])
sheet.append(['第二项', 5])
sheet.append(['合计', '=SUM(B2:B3)'])
sheet.add_image(SheetImage(str(picture)), 'D2')
chart = BarChart()
chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True)
sheet.add_chart(chart, 'D6')
sheet.print_area = 'A1:J20'
sheet.column_dimensions['A'].width = 24
sheet.column_dimensions['B'].width = 14
xlsx_path = root / 'budget.xlsx'
book.save(xlsx_path)
edited_book = load_workbook(xlsx_path)
edited_book.create_sheet('补充').append(['已修改工作表'])
edited_book.save(xlsx_path)
assert len(load_workbook(xlsx_path).sheetnames) == 2
print(json.dumps({'docx': str(doc_path), 'pptx': str(ppt_path), 'xlsx': str(xlsx_path)}, ensure_ascii=False))
